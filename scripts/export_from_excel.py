"""Export Debt.xlsx history into data/portfolio.json (plain, git-ignored).

Reads the cached (last-calculated) values, so open + save the workbook in
Excel first if you've edited it. Then run scripts/encrypt.mjs to produce the
encrypted file that is safe to commit.

    .venv/bin/python scripts/export_from_excel.py [path/to/Debt.xlsx]
"""

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Debt.xlsx"
OUT = ROOT / "data" / "portfolio.json"

# Asset id -> (display name, market value column, cumulative cost column or None).
# Order here fixes each asset's chart colour, so keep it stable.
ASSETS = [
    ("iwda", "IWDA / IBKR", "AJ", "AK"),
    ("sti", "STI ETF", "H", "G"),
    ("rsu", "TikTok RSU", "AQ", None),
    ("syfe_income", "SYFE Income+", "Q", "P"),
    ("reit", "Nikko AM REIT ETF", "AE", "REIT_COST"),
    ("cash", "Cash / SYFE Cash+", "U", None),
    ("hstech", "HSI-Tech ETF", "AT", "AR"),
    ("crypto", "Crypto", "AN", "AM"),
    ("eu_etf", "EU ETF", "AP", "AO"),
    ("fs", "Funding Societies", "AL", None),
]
# Cash is money, not an investment: its "cost" is always its value.
CASH_LIKE = {"cash"}
PASSIVE = [
    ("sti_div", "STI ETF dividends", "AW"),
    ("reit_div", "REIT dividends", "AZ"),
    ("php_reit_div", "Phillip REIT dividends", "BB"),
    ("bond_div", "SYFE Income+ payout", "AY"),
    ("cash_interest", "Cash returns", "BD"),
]


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def text(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:,.2f}"
    return str(v).strip()


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Investments & Savings"]

    months = {}
    nikko_cost = phillip_cost = 0.0
    for r in range(2, ws.max_row + 1):
        date = ws[f"A{r}"].value
        if not isinstance(date, datetime):
            continue
        key = date.strftime("%Y-%m")
        nikko_cost += num(ws[f"Z{r}"].value) or 0
        phillip_cost += num(ws[f"AA{r}"].value) or 0
        # Phillip REIT ETF (AD) was sold off in Nov 2024; drop its cost from then on.
        reit_cost = nikko_cost + (phillip_cost if key < "2024-11" else 0)

        values, invested = {}, {}
        for aid, _name, vcol, ccol in ASSETS:
            v = num(ws[f"{vcol}{r}"].value)
            if not v:
                continue
            values[aid] = round(v, 2)
            if ccol == "REIT_COST":
                invested[aid] = round(reit_cost, 2)
            elif ccol:
                c = num(ws[f"{ccol}{r}"].value)
                if c is not None:
                    invested[aid] = round(c, 2)

        passive = {}
        for pid, _name, col in PASSIVE:
            v = num(ws[f"{col}{r}"].value)
            if v and abs(v) > 0.005:
                passive[pid] = round(v, 2)

        notes = []
        for amt_col, label_col in (("V", "W"), ("X", "Y")):
            amt, label = text(ws[f"{amt_col}{r}"].value), text(ws[f"{label_col}{r}"].value)
            if amt and label and "\n" not in amt and "\n" not in label:
                notes.append(f"{label}: {amt}")
            elif amt and label:
                notes.append(" / ".join(
                    f"{l}: {a}" for a, l in zip(amt.split("\n"), label.split("\n"))
                ))
            elif amt or label:
                notes.append(label or amt)
        remark = ws[f"BC{r}"].value
        if num(remark) is not None and key >= "2025-07":
            notes.append(f"BTC held: {num(remark):g}")
        elif text(remark):
            notes.append(text(remark))

        # A later row for the same month (rare) wins.
        months[key] = {
            "month": key,
            "values": values,
            "invested": invested,
            "passive": passive,
            "note": "\n".join(n for n in notes if n),
        }

    ordered = sorted(months.values(), key=lambda m: m["month"])
    fill_gaps(ordered, [a[0] for a in ASSETS])
    budget = build_budget(wb["Money Transfer"])

    data = {
        "version": 1,
        "currency": "SGD",
        "updatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "assets": [
            {"id": aid, "name": name, "costTracked": ccol is not None and aid not in CASH_LIKE}
            for aid, name, _v, ccol in ASSETS
        ],
        "passiveTypes": [{"id": pid, "name": name} for pid, name, _c in PASSIVE],
        "months": ordered,
        "budgets": [budget],
    }
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(data, indent=1, ensure_ascii=False))
    print(f"Wrote {OUT.relative_to(ROOT)}: {len(data['months'])} months, "
          f"{data['months'][0]['month']} → {data['months'][-1]['month']}")


def fill_gaps(months, asset_ids, max_gap=3):
    """A few months were left blank in the sheet (e.g. 2022-09/10, 2024-03).
    Linearly interpolate an asset across a short gap and flag it as estimated."""
    for aid in asset_ids:
        known = [i for i, m in enumerate(months) if aid in m["values"]]
        for a, b in zip(known, known[1:]):
            if not 1 < b - a <= max_gap + 1:
                continue
            for i in range(a + 1, b):
                t = (i - a) / (b - a)
                for field in ("values", "invested"):
                    lo, hi = months[a][field].get(aid), months[b][field].get(aid)
                    if lo is not None and hi is not None:
                        months[i][field][aid] = round(lo + (hi - lo) * t, 2)
                months[i].setdefault("estimated", []).append(aid)


def build_budget(ws):
    """The plan in 'Money Transfer' column C (rows 2-26)."""
    salary = num(ws["C2"].value) or 0
    take_home = num(ws["C3"].value) or 0
    lines = lambda rows: [
        {"name": text(ws[f"B{r}"].value), "amount": round(num(ws[f"C{r}"].value) or 0, 2)}
        for r in rows
        if ws[f"B{r}"].value
    ]
    return {
        "month": datetime.now().strftime("%Y-%m"),
        "income": [{"name": "Salary", "amount": salary}],
        "deductions": [{"name": "CPF (employee) + CDAC", "amount": round(salary - take_home, 2)}],
        "expenses": lines(range(4, 15)),
        "investments": lines(range(15, 27)),
    }


if __name__ == "__main__":
    main()
